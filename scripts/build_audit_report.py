"""Build the academic audit supplement with the bundled python-docx runtime.

Run from the project root. Rendering/visual review is a separate release gate.
"""
from pathlib import Path
from docx import Document
from docx.shared import Cm, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

ROOT = Path(__file__).resolve().parents[1]
OUTPUT = ROOT / 'output/docx/BoskePOS_Auditoria_y_Plan_Proyecto_de_Titulo.docx'
doc = Document()
section = doc.sections[0]
section.page_width, section.page_height = Cm(21), Cm(29.7)
section.top_margin, section.bottom_margin = Cm(2), Cm(1.9)
section.left_margin, section.right_margin = Cm(2), Cm(2)
section.footer_distance = Cm(0.9)
for name in ['Normal', 'Title', 'Subtitle', 'Heading 1', 'Heading 2']:
    style = doc.styles[name]
    style.font.name = 'Arial'
    style.font.color.rgb = RGBColor(0, 0, 0)
doc.styles['Normal'].font.size = Pt(11)
doc.styles['Normal'].paragraph_format.space_after = Pt(8)
doc.styles['Normal'].paragraph_format.line_spacing = 1.12
doc.styles['Title'].font.size = Pt(26)
doc.styles['Title'].paragraph_format.space_after = Pt(16)
doc.styles['Heading 1'].font.size = Pt(18)
doc.styles['Heading 1'].paragraph_format.space_after = Pt(12)
doc.styles['Heading 2'].font.size = Pt(12)
doc.styles['Heading 2'].paragraph_format.space_before = Pt(12)
doc.styles['Heading 2'].paragraph_format.space_after = Pt(6)
for style in doc.styles:
    for border in list(style.element.iter(qn('w:pBdr'))):
        border.getparent().remove(border)
doc.core_properties.title = 'Auditoría y evolución de BoskePOS para Proyecto de Título'
doc.core_properties.subject = 'MVP v0.2 y revisión técnica v0.2.1'
doc.core_properties.author = 'BoskePOS'
doc.core_properties.keywords = 'auditoría, evidencia, trazabilidad, Proyecto de Título'
footer = section.footer.paragraphs[0]
footer.alignment = WD_ALIGN_PARAGRAPH.RIGHT
r = footer.add_run('BoskePOS  |  ')
r.font.size = Pt(9)
field = OxmlElement('w:fldSimple')
field.set(qn('w:instr'), 'PAGE')
footer._p.append(field)

def p(text, style=None):
    return doc.add_paragraph(text, style)

def h(text):
    doc.add_heading(text, 2)

def page(title):
    doc.add_page_break()
    doc.add_heading(title, 1)

def table(headers, rows, widths):
    t = doc.add_table(rows=1, cols=len(headers))
    t.alignment = WD_TABLE_ALIGNMENT.CENTER
    t.autofit = False
    for col, width in zip(t.columns, widths):
        col.width = Cm(width)
    for c, text, width in zip(t.rows[0].cells, headers, widths):
        c.text, c.width = text, Cm(width)
    for values in rows:
        cells = t.add_row().cells
        for c, text, width in zip(cells, values, widths):
            c.text, c.width = str(text), Cm(width)
    pr = t._tbl.tblPr
    borders = OxmlElement('w:tblBorders')
    for edge in ['top', 'left', 'bottom', 'right', 'insideH', 'insideV']:
        e = OxmlElement('w:' + edge)
        for key, val in [('val', 'single'), ('sz', '4'), ('color', 'D9D9D9')]:
            e.set(qn('w:' + key), val)
        borders.append(e)
    pr.append(borders)
    for ri, row in enumerate(t.rows):
        trpr = row._tr.get_or_add_trPr()
        trpr.append(OxmlElement('w:cantSplit'))
        if ri == 0:
            repeat = OxmlElement('w:tblHeader')
            trpr.append(repeat)
        for ci, c in enumerate(row.cells):
            c.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            tcpr = c._tc.get_or_add_tcPr()
            margins = OxmlElement('w:tcMar')
            for side in ['top', 'left', 'bottom', 'right']:
                m = OxmlElement('w:' + side)
                m.set(qn('w:w'), '90')
                m.set(qn('w:type'), 'dxa')
                margins.append(m)
            tcpr.append(margins)
            if ri == 0:
                shade = OxmlElement('w:shd')
                shade.set(qn('w:fill'), 'DCEAF4')
                tcpr.append(shade)
            for par in c.paragraphs:
                par.paragraph_format.space_after = Pt(0)
                par.paragraph_format.line_spacing = 1.05
                if ci == 0 and widths[0] < 3:
                    par.alignment = WD_ALIGN_PARAGRAPH.CENTER
                if headers[ci].endswith(' ms'):
                    par.alignment = WD_ALIGN_PARAGRAPH.RIGHT
                for run in par.runs:
                    run.font.size = Pt(10)
                    run.bold = ri == 0
    doc.add_paragraph().paragraph_format.space_after = Pt(0)
    return t

p('Auditoría y evolución de BoskePOS para Proyecto de Título', 'Title')
p('Evaluación técnica y plan de validación académica', 'Subtitle')
p('16 de septiembre de 2026  |  MVP v0.2 y revisión local v0.2.1')
h('Dictamen')
p('BoskePOS dispone de un núcleo funcional de ventas, compras e inventario que puede demostrarse y defenderse técnicamente. La revisión corrigió fallas de integridad, validación y experiencia de uso; la suite final aprobó 139 pruebas en PostgreSQL, incluidas pruebas con conexiones concurrentes.')
p('La recomendación es utilizar v0.2.1 como base de demostración académica y preparar un piloto supervisado. Antes de operar de forma habitual deben resolverse la valorización de compras anuladas, los reintentos de cobro y la recuperación de respaldos. La aprobación técnica no equivale a aceptación del negocio.')
h('Alcance de esta revisión')
p('Se revisaron el repositorio, la API, los servicios transaccionales, permisos, interfaz, dependencias y el informe de primera y segunda parte. Se ejecutó el sistema con una base de demostración aislada y se recorrió una venta ficticia desde búsqueda hasta anulación. El informe fuente de 21 páginas se contrastó con el comportamiento del código.')
table(['Dimensión', 'Resultado al cierre'], [
    ('Código', 'Correcciones locales v0.2.1 sobre la base v0.2; sin publicación ni despliegue productivo.'),
    ('Validación', '139 pruebas PostgreSQL aprobadas; lint y build aprobados; auditorías de dependencias sin avisos conocidos.'),
    ('Trabajo académico', 'Observaciones por página, trazabilidad, objetivos medibles, plan de piloto y roadmap.'),
    ('Pendiente', 'UAT real, métricas operativas, periféricos, valorización, idempotencia y continuidad ensayada.'),
], [3.1, 13.9])
h('Lectura recomendada')
p('Las páginas siguientes presentan correcciones, evidencia, revisión del informe, requisitos, versiones y validación. El detalle reproducible se conserva en docs/auditoria-mvp-v0.2.1.md y docs/evidence/2026-09-16. Los estados futuros se expresan como propuestas con criterios de salida.')

page('Correcciones implementadas')
p('Las 120 pruebas originales pasaban, pero no cubrían varios escenarios de estado desactualizado y entradas malformadas. Se añadieron 19 pruebas y se corrigieron los defectos reproducidos.')
table(['Área', 'Problema', 'Cambio aplicado'], [
    ('Inventario', 'Una edición de catálogo podía sobrescribir stock modificado por otra operación.', 'Guardar solo campos editados y mantener stock de solo lectura en Admin.'),
    ('Compras', 'Un borrador antiguo podía cambiar una compra ya confirmada.', 'Bloquear y releer el estado al editar o descartar el borrador.'),
    ('Admin', 'Cambios directos de estados o líneas evitaban las reglas de inventario.', 'Ventas, compras y sus líneas solo consultables en Django Admin.'),
    ('Ventas', 'PATCH sin ítems fallaba; ciertos montos excedían la capacidad de persistencia.', 'Conservar campos omitidos y validar capacidad monetaria antes de escribir.'),
    ('Consultas', 'Fechas e IDs inválidos llegaban al ORM; productos protegidos fallaban al borrarse.', 'Errores 400 de validación y 409 para eliminación incompatible con historia.'),
    ('Dashboard', 'Una venta suspendida se sumaba en su fecha de creación.', 'Usar fecha de cobro; fallback para registros antiguos sin completed_at.'),
    ('Comprobante', 'El vendedor no obtenía datos del negocio para imprimir.', 'Endpoint de lectura restringida y texto explícito de comprobante interno.'),
], [2.5, 7.2, 7.3])
h('Interfaz y operación')
p('El POS cancela consultas obsoletas y bloquea el cobro hasta obtener una cotización válida. Limpiar el carrito reinicia cliente y medio de pago. La búsqueda limita resultados a 20 y la navegación móvil conserva los permisos del usuario. Se añadió búsqueda por referencia en Kardex y proxy de imágenes para desarrollo.')
p('Se incorporaron límites por IP en login y refresh, guardas de configuración productiva, un health check sin datos sensibles, una demo aislada y un auditor de saldos. Las dependencias con avisos conocidos fueron actualizadas. El workflow CI queda preparado; no se ejecutó remotamente en esta revisión.')

page('Evidencia técnica y límites')
table(['Control', 'Resultado', 'Condición'], [
    ('Suite backend', '139 aprobadas', 'PostgreSQL 17.10; 99,137 segundos.'),
    ('Concurrencia', '2 aprobadas', 'Última unidad y doble anulación; conexiones reales.'),
    ('Frontend', 'Lint y build aprobados', 'No existe suite automática de componentes o E2E.'),
    ('Dependencias npm', '14 alertas iniciales a 0', 'Auditoría incluye herramientas de desarrollo.'),
    ('Dependencias Python', '12 avisos iniciales a 0', 'Tres paquetes afectados; requirements.txt auditado.'),
    ('Kardex demo', '16 productos sin diferencias', 'Conciliación lógica; no conteo físico.'),
    ('Despliegue', '2 advertencias HSTS', 'Perfil HTTPS temporal; no certifica producción.'),
], [4, 5.2, 7.8])
h('Medición local de la API')
p('Se midieron 100 solicitudes por escenario después de cinco de calentamiento, con un cliente secuencial, 16 productos ficticios, SQLite y runserver en localhost. La autenticación queda fuera del cronómetro. P95 se calcula por rango más cercano.')
table(['Escenario', 'Mediana ms', 'P95 ms', 'Máximo ms'], [
    ('Búsqueda', '3,99', '27,44', '29,92'),
    ('Cotización de una línea', '3,45', '23,13', '28,79'),
    ('Dashboard', '10,22', '31,77', '36,25'),
], [7.7, 3.1, 3.1, 3.1])
p('Estas cifras son una línea base de desarrollo. No miden confirmación persistida de venta, tiempo del cajero, renderizado del navegador, carga representativa en PostgreSQL ni disponibilidad de 99,5%.')
h('Recorrido de navegador')
p('Se probó ingreso, búsqueda, descuento inválido, suspensión, recuperación, cobro, comprobante, anulación y navegación móvil. En la venta ficticia 1, el producto AB001 pasó de 24 a 23 unidades y volvió a 24 al anular; el Kardex conserva ambos movimientos. Total 1.990, efectivo 2.000 y vuelto 10.')
p('Entorno verificado: Windows, Python 3.13.2, Django 6.0.8, DRF 3.17.2, sqlparse 0.6.0, Node 25.2.1 y npm 11.6.2. La ejecución completa usa PostgreSQL; la demo usa una base independiente. No se midió cobertura de código ni se realizó un pentest.')

page('Mejoras del informe académico')
p('Las ubicaciones corresponden a las 21 páginas del informe fuente renderizado en Word. La revisión detallada y el texto de reemplazo se encuentran en docs/proyecto-titulo/revision-informe.md.')
table(['Página', 'Corrección necesaria'], [
    ('2', 'Sustituir el índice manual por una tabla de contenido con estilos y páginas; separar XII y XIII.'),
    ('3 a 5', 'Sustentar diagnóstico con fechas, participantes, instrumentos y actas. Atribuir las percepciones sobre UVO.'),
    ('5 y 6', 'Marcar importación masiva pendiente; aclarar cantidades enteras. Compras ya implementadas, con límite de valorización.'),
    ('9 y 11', 'Actualizar estado v0.2/v0.2.1 y Django 6.0.8. Incluir system_settings; distinguir dashboard de reports.'),
    ('12 a 16', 'Revisar semántica del proceso; corregir diagrama que rebasa página 13; ampliar DER y topología y usar modelos reales.'),
    ('16 y 17', 'Separar metas KPI/SLA de resultados. Usar estados de prueba inequívocos e incluir HU-12.'),
    ('18 a 21', 'Fechar cronograma, precisar fuentes y costos. Corregir permisos: inventario no lee clientes ni ventas.'),
], [2.2, 14.8])
h('Objetivo general propuesto')
p('Desarrollar y validar un sistema web de gestión de ventas e inventario para una sucursal piloto de Ferretería El Bosque, que centralice los datos operativos y permita evaluar la reducción de fricción en caja mediante pruebas de integridad, tareas con usuarios y mediciones comparables con el procedimiento vigente.')
h('Argumento para la defensa')
p('El aporte consiste en adaptar un proceso concreto y comprobar su integridad, trazabilidad y facilidad de uso. Las apps Django forman un monolito modular con API y frontend separado; no son microservicios. No es necesario presentar el producto como un ERP completo ni afirmar superioridad general sobre otra plataforma.')
p('Para evaluar impacto, se propone comparar la mediana de tareas equivalentes antes y después, con una meta inicial de reducción de 20% sujeta a acuerdo. La cifra es una hipótesis de validación, no un beneficio ya demostrado. Los diagramas técnicos actualizados están en docs/architecture.md.')

page('Trazabilidad funcional')
p('Verificación técnica significa código y escenarios probados. La aceptación funcional debe registrarla una persona del rol correspondiente. La matriz completa enlaza métodos reales en docs/proyecto-titulo/trazabilidad.md.')
table(['Historia', 'Capacidad', 'Estado al corte'], [
    ('HU-01', 'Búsqueda y código de barras', 'Verificada; lector físico pendiente.'),
    ('HU-02', 'Venta y pago simple', 'Verificada; idempotencia pendiente.'),
    ('HU-03', 'Ajustes y Kardex', 'Verificada; conteo físico pendiente.'),
    ('HU-04', 'Dashboard', 'Verificado por período y rol; no BI.'),
    ('HU-05', 'Categorías', 'CRUD y validaciones verificados.'),
    ('HU-06', 'Clientes', 'CRUD, búsqueda y RUT verificados.'),
    ('HU-07', 'Historial de ventas', 'Paginación y filtros verificados.'),
    ('HU-08', 'Operación rápida y volumen', 'Parcial; cantidades fraccionarias pendientes.'),
    ('HU-09', 'Importación y exportación', 'Pendiente; propuesta v0.4.'),
    ('HU-10', 'Roles y permisos', 'Verificados; gestión de usuarios en Django Admin.'),
    ('HU-11', 'Anulación de ventas', 'Reposición única verificada; sin documento tributario.'),
    ('HU-12', 'Proveedores y compras', 'Recepción y reversa física verificadas; costo pendiente.'),
], [2.3, 6.1, 8.6])
h('Requisitos no funcionales')
p('La integridad dispone de evidencia de rollback y concurrencia, pero necesita ampliar reintentos y valorización. Los controles de roles y dependencias sustentan una revisión parcial de seguridad. El logout invalida el refresh; el access token sigue vigente hasta su expiración.')
p('La usabilidad y accesibilidad tienen un recorrido técnico inicial. Deben medirse con usuarios, teclado, lector de pantalla y dispositivos del negocio. El health check prepara el monitoreo; no demuestra disponibilidad sostenida. La mantenibilidad cuenta con servicios, pruebas y CI preparado, sin una ejecución remota acreditada.')
h('Regla de cierre')
p('Cada requisito debe enlazar implementación, prueba identificada, fecha, ambiente, resultado y limitación. Usar aprobado, fallido, bloqueado o no ejecutado. Evitar estados ambiguos como aprobado según release y porcentajes de cobertura sin medición.')

page('Versiones propuestas del MVP')
p('La secuencia prioriza confiabilidad operativa antes de integración tributaria. Las versiones posteriores a v0.2.1 son propuestas; su alcance y calendario deben ajustarse a la disponibilidad del equipo y del negocio.')
table(['Versión', 'Contenido', 'Condición de salida'], [
    ('v0.2.1', 'Estabilización realizada', '139 pruebas PostgreSQL, lint/build y recorrido técnico.'),
    ('v0.3', 'Idempotencia, valorización, redondeo CLP, caja y continuidad', 'Sin duplicados por reintento; caja conciliada; restauración medida.'),
    ('v0.4', 'Importación CSV, fracciones por unidad y accesibilidad', 'Vista previa y lote seguro; pruebas con catálogo y periféricos reales.'),
    ('v0.5', 'Devoluciones parciales, pagos múltiples y reportes', 'Pagos y devoluciones conciliados; costo histórico para márgenes.'),
    ('v0.6', 'Integración tributaria acotada', 'Alternativa y requisitos verificados; pruebas y autorización aplicable.'),
    ('v1.0', 'Operación aceptada', 'Piloto aprobado, soporte, manuales y continuidad demostrada.'),
], [2.2, 7.1, 7.7])
h('Aceptación prioritaria de v0 3')
p('Idempotencia: la misma clave y payload recuperan la misma venta; reutilizar la clave con otro payload produce conflicto. Probar solicitud simultánea y pérdida de respuesta después del commit. El bloqueo del botón por sí solo no cumple este criterio.')
p('Valorización: acordar el tratamiento de una compra anulada después de otras compras o ventas. Registrar costos históricos y conciliar cantidad y valor. No restaurar ciegamente un promedio anterior porque puede invalidar operaciones posteriores.')
p('Caja y continuidad: registrar apertura, fondo, ingresos, egresos y cierre por responsable; calcular diferencias desde eventos. Ensayar restauración en otra base, medir RPO/RTO y monitorizar un staging con PostgreSQL y HTTPS.')
h('Plan de iteraciones')
p('Primero revisar v0.2.1 y acordar reglas de dinero/costos; después implementar idempotencia y valorización, caja y staging; finalmente ejecutar UAT y corregir. Planificar v0.4 según lo aprendido. Multibodega, offline y BI quedan fuera del alcance base. Backlog y estimaciones orientativas: docs/roadmap-mvp.md.')

page('Validación con usuarios y métricas')
p('Se propone una evaluación exploratoria con propietario, dos vendedores y un encargado de inventario, si existen esos roles disponibles. Cada participante realiza cinco repeticiones por tarea de su rol, después de una familiarización uniforme. Informar la muestra real y sus límites.')
p('Comparar tareas equivalentes en el procedimiento vigente y BoskePOS, con datos ficticios y el mismo entorno. Alternar el orden cuando sea viable. Registrar tiempo, pasos, exactitud, ayuda y errores, además de versión, catálogo, red y dispositivo.')
table(['Indicador', 'Método y meta propuesta'], [
    ('Tiempo de tarea', '(Mediana base − mediana nueva) / mediana base × 100; mejora ≥20%, a acordar.'),
    ('Búsqueda', 'P95 ≤1 s; medir HTTP y experiencia del navegador por separado.'),
    ('Confirmación', 'P95 ≤2 s hasta respuesta persistida; PostgreSQL y carga declarada.'),
    ('Stock físico', 'SKU exactos / SKU contados ×100; meta ≥98%; conteo real pendiente.'),
    ('Kardex', 'Productos sin diferencias / auditados ×100; meta 100%, sin escrituras concurrentes.'),
    ('Éxito técnico', 'Operaciones válidas sin error / intentos válidos ×100; meta ≥99%.'),
    ('Disponibilidad', 'Tiempo disponible / tiempo comprometido ×100; meta ≥99,5%; acordar horario.'),
    ('Recuperación', 'RTO ≤4 h y RPO ≤24 h; medir restauración y pérdida real en un ensayo.'),
], [4, 13])
h('Casos mínimos de aceptación')
p('Buscar con escáner; vender tres productos; rechazar stock o descuento inválido; suspender y recuperar; anular con reposición única; recibir compra y ajustar con motivo; rechazar acciones por rol; recuperar un envío incierto; restaurar respaldo; completar tareas por teclado. El envío incierto requiere implementar idempotencia antes de aprobarse.')
p('Para carga se proponen 1, 5 y 10 usuarios y catálogos ficticios de 1.000 y 10.000 productos, ajustados al volumen real. Registrar errores y recursos, además de latencia. El plan detallado de ejecución, reversión y viabilidad está en docs/proyecto-titulo/plan-validacion.md.')

page('Riesgos y preparación de la entrega')
h('Riesgos que condicionan el piloto')
p('Valorización y reintentos son prioridades altas. Hoy anular una compra compensa unidades sin revertir el costo promedio. Una respuesta de cobro perdida puede generar otra venta si se reintenta. También deben acordarse importes CLP y cantidades fraccionarias: la base guarda dinero con dos decimales y cantidades enteras.')
p('La operación requiere respaldo externo, restauración y monitoreo. Los límites por IP usan caché local y no se comparten automáticamente entre procesos. Faltan pruebas de periféricos y foco en modales. La reimpresión consulta datos actuales del negocio; no conserva una instantánea de identidad comercial.')
h('Evidencia para la evaluación del ramo')
p('Completar actas e instrumentos de diagnóstico, línea base de tareas, rúbrica institucional, resultados UAT y aceptación del piloto. El presupuesto debe contemplar horas por tarifa declarada, infraestructura, respaldo, soporte y contingencia. Sin cotizaciones y horas verificables no corresponde declarar un ROI.')
p('En la defensa, demostrar una venta y su anulación, explicar una validación fallida y presentar la prueba de última unidad en PostgreSQL. Diferenciar consistencia del Kardex de exactitud física. Explicar qué está implementado, qué se probó y qué permanece planificado.')
h('Fuentes técnicas')
sources = [
    ('Django Software Foundation', 'QuerySet API y select_for_update', 'https://docs.djangoproject.com/en/6.0/ref/models/querysets/#select-for-update'),
    ('Django Software Foundation', 'Deployment checklist', 'https://docs.djangoproject.com/en/6.0/howto/deployment/checklist/'),
    ('Django Software Foundation', 'Django 6.0.8 release notes', 'https://docs.djangoproject.com/en/6.0/releases/6.0.8/'),
    ('Django REST Framework', 'Release notes 3.17.2', 'https://www.django-rest-framework.org/community/release-notes/#3172'),
    ('Django REST Framework', 'Throttling', 'https://www.django-rest-framework.org/api-guide/throttling/'),
]
for author, title, url in sources:
    par = p(f'{author}. {title}. Consultado el 16 de septiembre de 2026.\n{url}')
    par.paragraph_format.space_after = Pt(7)
    for run in par.runs:
        run.font.size = Pt(9)
p('Base del análisis: código de BoskePOS desde 6dbb13c, cambios locales y registros de ejecución. Informe revisado: BoskePOS_Informe_Primera_y_Segunda_Parte.docx. Las fuentes técnicas explican mecanismos y parches; no prueban resultados del negocio.')

OUTPUT.parent.mkdir(parents=True, exist_ok=True)
doc.save(OUTPUT)
print(OUTPUT)
